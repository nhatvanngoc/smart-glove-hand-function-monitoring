"""Regression tests for clean proposal exports; no firmware or hardware tests.

Run: python3 -m unittest discover -s tests -v
Requires python-docx, as does the exporter. Temporary DOCX files are not retained.
"""
import importlib.util
from pathlib import Path
import re
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("abc_builder", ROOT / "scripts/build_abc_review_docx.py")
builder = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(builder)
SOURCE = ROOT / "docs/DE_CUONG_A_B_C_BAN_TRONG_TAM_2026-10-20.md"


def text_of(document):
    return "\n".join(
        [p.text for p in document.paragraphs]
        + [c.text for t in document.tables for row in t.rows for c in row.cells]
    )


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.clean = Path(cls.temp.name) / "clean.docx"
        cls.small = Path(cls.temp.name) / "13pt.docx"
        cls.review = Path(cls.temp.name) / "review.docx"
        builder.build(str(SOURCE), str(cls.clean), no_meta=True, cover_version="2026-10-20")
        builder.build(str(SOURCE), str(cls.small), no_meta=True, cover_version="2026-10-20", size=13, spacing=1.15)
        builder.build(str(SOURCE), str(cls.review))

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_clean_check_does_not_require_metadata_date(self):
        for path in (self.clean, self.small):
            with self.subTest(path=path.name):
                failures, info = builder.verify(str(path), str(SOURCE), no_meta=True)
                self.assertEqual([], failures)
                self.assertEqual(7, info["tables"])
                self.assertEqual("không chèn metadata", info["ngay_dung"])

    def test_review_mode_retains_metadata_and_valid_date(self):
        failures, info = builder.verify(str(self.review), str(SOURCE))
        self.assertEqual([], failures)
        self.assertEqual(8, info["tables"])
        self.assertRegex(info["ngay_dung"], r"^\d{4}-\d{2}-\d{2}$")
        self.assertEqual("Nguồn Markdown", Document(self.review).tables[0].cell(0, 0).text)

    def test_clean_checker_rejects_metadata_even_if_body_is_clean(self):
        failures, _ = builder.verify(str(self.review), str(SOURCE), no_meta=True)
        self.assertTrue(any("thông tin vận hành" in failure for failure in failures))

    def test_cover_contents_body_and_footer(self):
        doc = Document(self.clean)
        self.assertEqual("ĐỀ CƯƠNG DỰ ÁN KHOA HỌC KỸ THUẬT", doc.paragraphs[0].text)
        heads = {p.text: p for p in doc.paragraphs if p.style.name.startswith("Heading")}
        self.assertTrue(heads["MỤC LỤC"].paragraph_format.page_break_before)
        self.assertTrue(heads["A. VẤN ĐỀ NGHIÊN CỨU"].paragraph_format.page_break_before)
        self.assertTrue(heads["PHỤ LỤC — DANH MỤC HÌNH CẦN BỔ SUNG"].paragraph_format.page_break_before)
        self.assertTrue(doc.sections[0].different_first_page_header_footer)
        self.assertEqual("Phiên bản: 2026-10-20", doc.sections[0].first_page_footer.paragraphs[0].text)
        self.assertNotIn("Phiên bản:", text_of(doc))
        self.assertEqual(10, doc.sections[0].first_page_footer.paragraphs[0].runs[0].font.size.pt)

    def test_size_changes_real_body_heading_and_table_runs(self):
        for path, size in ((self.clean, 14), (self.small, 13)):
            with self.subTest(size=size):
                doc = Document(path)
                body = next(p for p in doc.paragraphs if p.text.startswith("Đột quỵ là"))
                heading = next(p for p in doc.paragraphs
                               if p.text == "A. VẤN ĐỀ NGHIÊN CỨU" and p.style.name.startswith("Heading"))
                self.assertTrue(all(r.font.size.pt == size for r in body.runs if r.text))
                self.assertTrue(all(r.font.size.pt == size for r in heading.runs if r.text))
                self.assertTrue(all(r.font.size.pt == size - 2
                                    for t in doc.tables for row in t.rows for c in row.cells
                                    for p in c.paragraphs for r in p.runs if r.text))
        # Building 13 pt must not mutate the module's default for subsequent exports.
        self.assertEqual(14, builder.SIZE)
        default = next(p for p in Document(self.review).paragraphs if p.text.startswith("Đột quỵ là"))
        self.assertEqual(14, default.runs[0].font.size.pt)

    def test_tables_repeat_headers(self):
        for table in Document(self.clean).tables:
            self.assertIsNotNone(table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")))

    def test_inline_rendering_preserves_paths_and_font_size(self):
        doc = Document()
        p = doc.add_paragraph()
        builder.emit_inline(p, "**Đậm** *nghiêng* `mã` [tài liệu](../docs/test.md)", size=13)
        self.assertEqual("Đậm nghiêng mã tài liệu (docs/test.md)", p.text)
        self.assertEqual(13, p.runs[0].font.size.pt)
        self.assertTrue(p.runs[0].bold)

    def test_empty_document_reports_failures_without_index_error(self):
        blank = Path(self.temp.name) / "blank.docx"
        Document().save(blank)
        failures, _ = builder.verify(str(blank), str(SOURCE), no_meta=True)
        self.assertTrue(any("thiếu bảng" in failure for failure in failures))

    def test_invalid_version_and_size_are_rejected(self):
        for kwargs in ({"no_meta": True, "cover_version": "2026-99-99"},
                       {"no_meta": True, "size": 0}, {"cover_version": "2026-10-20"}):
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                builder.build(str(SOURCE), str(Path(self.temp.name) / "invalid.docx"), **kwargs)


class AntiSlopContentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source = SOURCE.read_text(encoding="utf-8")
        cls.body = cls.source[cls.source.index("## A."):cls.source.index("## TÀI LIỆU THAM KHẢO")]

    def test_removed_templates_and_placeholders(self):
        for marker in ("**Cơ sở lý thuyết.**", "**Thiết kế.**", "Kiểm tra và kiểm thử tương ứng:",
                       "**Đã làm:**", "**Chưa làm:**", "⟨HÌNH", "Nguồn Markdown", "Quy tắc bằng chứng"):
            with self.subTest(marker=marker):
                self.assertNotIn(marker, self.source)
        self.assertEqual(1, self.body.count("chưa có biên bản đo"))

    def test_seven_figure_references_and_appendix_entries(self):
        for n in range(1, 8):
            self.assertEqual(1, len(re.findall(rf"\bHình {n}\b", self.body)))
            self.assertIn(f"| Hình {n} |", self.source)

    def test_pairing_table_and_method_column(self):
        for b, c in zip(range(3, 9), range(1, 7)):
            self.assertIn(f"| B.{b} | C.{c} |", self.body)
        self.assertIn("| Hạng mục | Phương pháp đo |", self.body)
        self.assertIn("| Tham số | Giá trị đang dùng | Cách kiểm chứng", self.body)

    def test_key_values_scope_and_status_are_preserved(self):
        for value in ("18 dB", "R²", "0,90", "90%", "85%", "5%", "8%", "18%", "60",
                      "20 ms", "80–120", "300 µs", "180 mV", "860", "0,998", "0,002", "0,20",
                      "55–80 s", "40–65 s", "28 s", "2.838 mV", "1,44–1,92 V", "50–135 mV",
                      "35", "24", "22 Hz", "120 µs", "1,0–1,3", "5,5"):
            self.assertIn(value, self.body)
        self.assertIn("| 9 | ngàm xoay 2 bản lề kép; đã in | chưa đo — mốc 2 | chưa đo — mốc 2 |", self.body)
        self.assertIn("tình trạng của bộ chuyển đổi ADS1115 chưa được xác nhận", self.body)
        self.assertIn("Đề tài không thử trên người bệnh", self.body)
        self.assertNotRegex(self.body, r"\b22\s*(?:s\b|giây)|\b5\s*N\b|khả thi nhất|10 phiên")

    def test_reference_entries_are_numbered_one_through_fifteen(self):
        refs = self.source.split("## TÀI LIỆU THAM KHẢO", 1)[1].split("## PHỤ LỤC", 1)[0]
        self.assertEqual(list(range(1, 16)), [int(n) for n in re.findall(r"^\[(\d+)\]", refs, re.M)])


if __name__ == "__main__":
    unittest.main()
